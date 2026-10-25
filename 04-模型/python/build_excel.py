from pathlib import Path
from openpyxl import Workbook

SHEETS = ["总表"] + [f"K{i:02d}" for i in range(1, 9)] + ["情景", "敏感性", "数据字典"]


def build_workbook(path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        ws = wb.create_sheet(name)
        ws["A1"] = name

    summary = wb["总表"]
    summary["A1"] = "王朝资产负债表"
    for col, value in enumerate(["项目", "金额", "单位", "来源"], start=1):
        summary.cell(row=3, column=col, value=value)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main() -> None:
    project_root = Path(__file__).resolve().parents[2]
    out = project_root / "04-模型/excel/王朝资产负债表.xlsx"
    build_workbook(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
