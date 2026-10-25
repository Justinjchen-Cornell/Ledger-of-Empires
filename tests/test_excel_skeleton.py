import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "04-模型/python"))
import build_excel

def test_build_workbook(tmp_path):
    path = tmp_path / "test.xlsx"
    build_excel.build_workbook(path)
    wb = load_workbook(path)
    assert "总表" in wb.sheetnames
    assert "K01" in wb.sheetnames
    assert "K08" in wb.sheetnames
    assert "情景" in wb.sheetnames
    assert wb["总表"]["A1"].value == "王朝资产负债表"
